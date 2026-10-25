"""Resume text extraction for PDF and DOCX.

Never raises on bad files: returns a status (ok / partial / failed) plus
warnings so the UI can flag problems instead of silently continuing.
"""

import io
import logging
from dataclasses import dataclass, field

from app.models import ParseStatus

log = logging.getLogger(__name__)

MIN_USEFUL_CHARS = 200  # less than this from a whole resume is suspicious
ALLOWED_EXTENSIONS = {".pdf", ".docx"}


@dataclass
class ParseResult:
    text: str = ""
    status: ParseStatus = ParseStatus.failed
    warnings: list[str] = field(default_factory=list)


def sniff_type(filename: str, data: bytes) -> str | None:
    """Return 'pdf' / 'docx' if the extension and file signature agree, else None."""
    name = filename.lower()
    if name.endswith(".pdf") and data[:5] == b"%PDF-":
        return "pdf"
    if name.endswith(".docx") and data[:4] == b"PK\x03\x04":
        return "docx"
    return None


def parse_resume(filename: str, data: bytes) -> ParseResult:
    kind = sniff_type(filename, data)
    if kind is None:
        return ParseResult(warnings=["File type not recognised. Upload a PDF or DOCX file."])
    try:
        result = _parse_pdf(data) if kind == "pdf" else _parse_docx(data)
    except Exception as exc:  # corrupt or encrypted files, odd encodings, etc.
        log.warning("Resume parsing failed (%s): %s", kind, type(exc).__name__)
        return ParseResult(warnings=[f"Could not read the {kind.upper()} file ({type(exc).__name__})."])

    result.text = _clean(result.text)
    if not result.text:
        result.status = ParseStatus.failed
        result.warnings.append(
            "No text could be extracted. The file may be a scanned image; OCR is not supported in v1."
        )
    elif len(result.text) < MIN_USEFUL_CHARS:
        result.status = ParseStatus.partial
        result.warnings.append(
            f"Only {len(result.text)} characters extracted; the resume may be incomplete."
        )
    elif result.warnings:
        result.status = ParseStatus.partial
    else:
        result.status = ParseStatus.ok
    return result


def _parse_pdf(data: bytes) -> ParseResult:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    warnings: list[str] = []
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            return ParseResult(warnings=["PDF is password protected."])
    pages = []
    empty = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            txt = page.extract_text() or ""
        except Exception:
            txt = ""
        if txt.strip():
            pages.append(txt)
        else:
            empty.append(i)
    if empty and pages:
        warnings.append(
            f"No text on page(s) {', '.join(map(str, empty))} of {len(reader.pages)}; "
            "they may be images."
        )
    return ParseResult(text="\n\n".join(pages), warnings=warnings)


def _parse_docx(data: bytes) -> ParseResult:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(dict.fromkeys(cells)))  # merged cells repeat text
    return ParseResult(text="\n".join(parts))


def _clean(text: str) -> str:
    lines = [" ".join(line.split()) for line in text.replace("\x00", "").splitlines()]
    out: list[str] = []
    for line in lines:
        if line or (out and out[-1]):
            out.append(line)
    return "\n".join(out).strip()
